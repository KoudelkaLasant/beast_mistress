import pygame
import os
import sys
import time
import contextlib
import copy
import random
from pathlib import Path
from configparser import ConfigParser
import ctypes
import pandas as pd
import openpyxl
import json

tile_width = 32
tile_height = 32
beastTeamMaximum = 3
	
class ExternalDataReader():
	def setup():
		ExternalDataReader.mapData = {}
	def fetch(relative_path):
		base_path = ""
		if getattr(sys, "frozen", False):
			base_path = Path(sys._MEIPASS)
		else:
			base_path = Path(__file__).resolve().parent
		return str(base_path / relative_path)
	def readMapData(map_name):
		source = ExternalDataReader.fetch(os.path.join("assets/maps", map_name + ".xlsx"))
		return pd.read_excel(source, sheet_name=None, header=None)
	def readCutscenes():
		source = ExternalDataReader.fetch(os.path.join("assets/data/cutscenes.xlsx"))
		workbook = openpyxl.load_workbook(source, data_only=True)
		results = {}
		for sheet in workbook.worksheets:
			headers = [cell.value for cell in sheet[1]]
			rows = []
			for row in sheet.iter_rows(min_row=2, values_only=True):
				rows.append(dict(zip(headers, row)))
			results[sheet.title] = rows
		return results
	def readDefault(fileName):
		source = ExternalDataReader.fetch(os.path.join(f"assets/data/{fileName}.xlsx"))
		workbook = openpyxl.load_workbook(source, data_only=True)
		sheet = workbook.active
		headers = [cell.value for cell in sheet[1]]
		skills = {}
		for row in sheet.iter_rows(min_row=2, values_only=True):
			entry = dict(zip(headers, row))
			skills[entry["UniqueID"]] = entry
		return skills
	def readBeasts():
		source = ExternalDataReader.fetch(os.path.join(f"assets/data/beasts.xlsx"))
		workbook = openpyxl.load_workbook(source, data_only=True)
		worksheet = workbook.worksheets[0]
		headers = [cell.value for cell in worksheet[1]]
		speciesStats = {}
		for row in worksheet.iter_rows(min_row=2, values_only=True):
			entry = dict(zip(headers, row))
			speciesStats[entry["Species Name"]] = entry
		worksheet = workbook.worksheets[1]
		headers = [cell.value for cell in worksheet[1]]
		beastLookupByName = {}
		for row in worksheet.iter_rows(min_row=2, values_only=True):
			entry = dict(zip(headers, row))
			speciesName = entry["Species Name"]
			del entry["Species Name"]
			for element in entry.keys():
				beastName = entry[element]
				beastElements = element.split("/")
				beastSpecies = speciesName
				beastLookupByName[beastName] = {"name" : beastName, "elements" : beastElements, "species" : beastSpecies}
		return speciesStats, beastLookupByName
		
		

	